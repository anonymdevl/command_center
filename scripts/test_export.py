"""The exports, tested without a site.

    python3 scripts/test_export.py

Frappe is stood in for, so what is under test is api/export.py's own work: that a
workbook opens, has one sheet per table with unique, legal names, carries the header
block, and gives Excel numbers rather than text for figures; that the PDF template
fills without a stray brace and escapes the question; and that the model's markdown is
sanitised before it reaches the PDF. Frappe's own PDF engine and md_to_html are the
site's, and are exercised there.
"""

from __future__ import annotations

import io
import os
import sys
import types

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

frappe = types.ModuleType("frappe")
frappe.local = types.SimpleNamespace(response=types.SimpleNamespace())
frappe.session = types.SimpleNamespace(user="kofi@gigmann.test")
frappe.whitelist = lambda **k: (lambda f: f)
frappe.parse_json = lambda s: __import__("json").loads(s)
frappe.db = types.SimpleNamespace(get_value=lambda *a, **k: "Gigmann Medical Supplies Ltd")
frappe.utils = types.SimpleNamespace(
    today=lambda: "2026-10-18", get_fullname=lambda u: "Kofi Mensah",
    now_datetime=lambda: None, format_datetime=lambda d, f=None: "18 Oct 2026, 09:00",
    escape_html=lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
    md_to_html=lambda md: f"<p>{md}</p><script>alert(1)</script>")
sys.modules["frappe"] = frappe
sys.modules["frappe.utils"] = frappe.utils
html_utils = types.ModuleType("frappe.utils.html_utils")
html_utils.sanitize_html = lambda h: h.replace("<script>alert(1)</script>", "")
sys.modules["frappe.utils.html_utils"] = html_utils
pdf = types.ModuleType("frappe.utils.pdf")
pdf.get_pdf = lambda html, options=None: html.encode()
sys.modules["frappe.utils.pdf"] = pdf
biz = types.ModuleType("command_center.api.businesses")
biz.require_manager = lambda: None
sys.modules["command_center.api.businesses"] = biz
www = types.ModuleType("command_center.www.command_center")
www._as_of = lambda: {"date": "2025-08-19"}
sys.modules["command_center.www.command_center"] = www

from openpyxl import load_workbook  # noqa: E402

from command_center.api import export  # noqa: E402

passed = failed = 0


def check(ok, what):
    global passed, failed
    passed += bool(ok)
    failed += not ok
    if not ok:
        print("FAIL:", what)


# Tables to Excel
export.tables_xlsx("Who owes us", [
    {"name": "Largest balances", "columns": ["Customer", "Owed (GHS)", "Share"],
     "rows": [["UBUNTU", "GHS 636,675.00", "6.9%"], ["A.S Hospitex", "541,841", "5.9%"]]},
    {"name": "Largest balances", "columns": ["A"], "rows": [["x"]]},
    {"name": "Bad/name:[x]?", "columns": ["A"], "rows": [["y"]]},
], "GIGMANN")
r = frappe.local.response
wb = load_workbook(io.BytesIO(r.filecontent))
check(r.type == "binary" and r.filename.endswith(".xlsx"), "an xlsx download")
check(len(wb.sheetnames) == 3 and len(set(wb.sheetnames)) == 3, "one sheet per table, names unique")
check(all(not set(n) & set("[]*?/\\:") for n in wb.sheetnames), "sheet names are legal")
ws = wb[wb.sheetnames[0]]
check(ws["A1"].value == "Who owes us" and "figures as at 2025-08-19" in ws["A2"].value,
      "title and the as-of header block")
check(ws["B5"].value == 636675.0 and abs(ws["C5"].value - 0.069) < 1e-9,
      "figures arrive as numbers, percentages as fractions")
check(ws["A5"].value == "UBUNTU", "text stays text")

# Answer to PDF
export.answer_pdf("Who owes <us>?", "**GHS 9.18M** owed", ["get_figures({})"], "GIGMANN",
                  "gemini-2.5-flash")
r = frappe.local.response
html = r.filecontent.decode()
check(r.type == "pdf" and r.filename.endswith(".pdf"), "a pdf download")
check("Who owes &lt;us&gt;?" in html, "the question is escaped")
check("<script>" not in html, "the model's markdown is sanitised")
check("Figures as at 2025-08-19" in html and "Kofi Mensah" in html, "the header says as-of and who")
check("get_figures" not in html and "What it checked" not in html,
      "what it checked stays on the screen and never reaches the PDF")
check("cannot change anything" in html, "the PDF states the boundary")

# Figures and charts are drawn on paper as on the screen, never dropped or printed raw.
md = ('Lead.\n\n```figures\n[{"label": "Owed", "value": "GHS 9.18M", "note": "806", "tone": "bad"}]\n```\n\n'
      '```chart\n{"type": "bar", "title": "Spend", "unit": "GHS", "labels": ["A", "B"], '
      '"series": [{"name": "Spend", "values": [10, 5]}]}\n```\n\n'
      '```chart\n{"type": "line", "labels": ["Jan", "Feb"], "series": [{"name": "S", "values": [1, 2]}]}\n```\n\n'
      '```chart\n{"type": "donut", "labels": ["X", "Y"], "series": [{"name": "S", "values": [3, 1]}]}\n```\n\n'
      '```chart\nnot json\n```\n\n```followups\n["Next?"]\n```\n\n'
      '| Who | Owed |\n|---|---|\n| UBUNTU | 636,675 |\n\n- **Recommendation:** Call them.')
export.answer_pdf("Report", md, [], "GIGMANN", "m")
doc = frappe.local.response.filecontent.decode()
check('class="figs"' in doc and "GHS 9.18M" in doc and "#c2392f" in doc,
      "a figures block becomes cards, in the tone's colour")
check('class="bars"' in doc and "Spend" in doc, "a bar chart is drawn as bars")
check(doc.count("<svg") == 2 and "polyline" in doc and "<path" in doc, "line and donut are drawn as SVG")
check("not json" not in doc and "```" not in doc and "CCBLOCK" not in doc,
      "a block that does not parse is dropped, not printed raw")
check("Next?" not in doc, "follow-up suggestions are left off paper")
_pol = export._polish('<p>Lead</p><p>More</p><table><tr><td>UBUNTU</td><td>636,675</td><td>GHS 9.18M</td>'
                      '<td>73.8%</td></tr></table><ul><li><strong>Recommendation:</strong> Call.</li></ul>')
check('<td class="n">636,675</td>' in _pol and '<td class="n">GHS 9.18M</td>' in _pol
      and '<td class="n">73.8%</td>' in _pol and "<td>UBUNTU</td>" in _pol,
      "numbers in tables line up on the right, words do not")
check('<li class="rec">' in _pol and _pol.count('class="lead"') == 1,
      "recommendations stand out; only the first paragraph is the lead")
check("<script" not in export.pdf_blocks.figures([{"label": "<script>x</script>", "value": "1"}]),
      "a card's text is escaped")

# _cell
check(export._cell("GHS 1,234.50") == 1234.5 and export._cell("12%") == 0.12, "figure parsing")
check(export._cell("Q3 2025") == "Q3 2025" and export._cell(None) == "", "non-figures untouched")

print(f"export: {passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
