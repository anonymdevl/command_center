"""Taking an answer, a table or the records behind a figure out of the platform.

Three downloads, all management-only, all reads -- nothing here creates a record.
Files are streamed back as the response; no File document is written to the site.

    answer_pdf     an Ask AI answer as a PDF: Gigmann's header, the question, the
                   answer, what it checked, as-of date, who generated it and when
    tables_xlsx    one or more tables (from an answer or a screen) as an .xlsx, one
                   sheet each
    records_xlsx   every record behind a figure -- not just the top rows the drawer
                   shows -- up to 20,000

Screens are printed by the browser (print styles in shell.py), so the PDF is exactly
what the manager is looking at. The answer PDF is built here because it should not
depend on which screen happens to be open.
"""

from __future__ import annotations

import io
import re

import frappe
from frappe.utils import escape_html

from command_center.api.businesses import require_manager

EXPORT_CAP = 20000
MAX_TABLES = 20
MAX_ROWS = 5000


def _name(title: str, ext: str) -> str:
    stem = re.sub(r"[^\w\- ]+", "", (title or "Command Center")).strip()[:60] or "Command Center"
    return f"{stem} {frappe.utils.today()}.{ext}"


def _business(business_code) -> str:
    if business_code and business_code != "__all__":
        return frappe.db.get_value("Connected Business", business_code, "business_name") or business_code
    return "All businesses"


def _as_of() -> str:
    from command_center.www.command_center import _as_of as boot_as_of
    d = boot_as_of() or {}
    return d.get("date") or "not loaded"


@frappe.whitelist(methods=["POST"])
def answer_pdf(question: str, answer: str, checked=None, business_code: str | None = None,
               model: str | None = None):
    """The answer as a PDF. The answer text is the one the analyst returned."""
    require_manager()
    from frappe.utils.html_utils import sanitize_html
    from frappe.utils.pdf import get_pdf

    checked = frappe.parse_json(checked) if isinstance(checked, str) else (checked or [])
    # The answer is model-written markdown: rendered, then sanitised, never trusted.
    body = sanitize_html(frappe.utils.md_to_html(_blocks_to_tables(answer or "")))
    who = frappe.utils.get_fullname(frappe.session.user)
    when = frappe.utils.format_datetime(frappe.utils.now_datetime(), "d MMM yyyy, HH:mm")
    checks = "".join(f"<li>{escape_html(str(c))}</li>" for c in checked[:40])
    html = PDF_TEMPLATE.format(
        business=escape_html(_business(business_code)),
        question=escape_html(question or ""),
        body=body,
        checks=(f"<h3>What it checked</h3><ul class='checks'>{checks}</ul>" if checks else ""),
        as_of=escape_html(_as_of()),
        who=escape_html(who),
        when=escape_html(when),
        model=escape_html(model or "the analyst"),
    )
    frappe.local.response.filename = _name(question[:60] if question else "Answer", "pdf")
    frappe.local.response.filecontent = get_pdf(html, {"page-size": "A4", "margin-top": "16mm",
                                                       "margin-bottom": "16mm",
                                                       "margin-left": "14mm",
                                                       "margin-right": "14mm"})
    frappe.local.response.type = "pdf"


@frappe.whitelist(methods=["POST"])
def tables_xlsx(title: str, tables, business_code: str | None = None):
    """Tables as an .xlsx, one sheet each. tables: [{name, columns: [..], rows: [[..]]}]"""
    require_manager()
    tables = frappe.parse_json(tables) if isinstance(tables, str) else (tables or [])
    sheets = []
    for i, t in enumerate(tables[:MAX_TABLES]):
        cols = [str(c) for c in (t.get("columns") or [])][:60]
        rows = [[_cell(v) for v in (r or [])][:60] for r in (t.get("rows") or [])[:MAX_ROWS]]
        sheets.append((t.get("name") or f"Table {i + 1}", cols, rows))
    _send_xlsx(title or "Tables", sheets, business_code)


@frappe.whitelist(methods=["POST"])
def records_xlsx(fact: str, filters=None, business_code: str | None = None, title: str | None = None):
    """Every record behind a figure, up to EXPORT_CAP, in the drawer's own columns."""
    require_manager()
    from command_center.api.facts import _records

    data = _records(fact, filters, business_code, limit=EXPORT_CAP, cap=EXPORT_CAP)
    cols = data.get("columns") or []
    header = [c.get("label") or c.get("key") for c in cols] + ["ERPNext link"]
    rows = [[_cell(r.get(c.get("key"))) for c in cols] + [r.get("open_url") or ""]
            for r in data.get("rows") or []]
    s = data.get("summary") or {}
    note = (f"{s.get('shown', len(rows))} of {s.get('count', len(rows))} records"
            + (" -- the export cap was reached" if s.get("count", 0) > len(rows) else ""))
    _send_xlsx(title or data.get("noun") or fact, [(title or "Records", header, rows)],
               business_code, note=note)


def _blocks_to_tables(md: str) -> str:
    """```figures and ```chart blocks become markdown tables on paper.

    The screen draws them (charts.js); a PDF engine cannot, and a PDF that silently
    dropped the key figures would be a worse document than the screen it came from.
    """
    import json

    def table(head, rows):
        out = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
        out += ["| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in rows]
        return "\n".join(out)

    def convert(m):
        lang, body = m.group(1).lower(), m.group(2)
        try:
            data = json.loads(body)
        except ValueError:
            return ""
        if lang == "figures" and isinstance(data, list):
            return table(["Figure", "Value", "Note"],
                         [[f.get("label", ""), f.get("value", ""), f.get("note", "")]
                          for f in data if isinstance(f, dict)])
        if lang == "chart" and isinstance(data, dict):
            series = [x for x in data.get("series") or [] if isinstance(x, dict)][:3]
            head = [""] + [x.get("name") or "Value" for x in series]
            rows = [[label] + [(x.get("values") or [])[i] if i < len(x.get("values") or []) else ""
                               for x in series]
                    for i, label in enumerate(data.get("labels") or [])]
            title = f"**{data.get('title')}**" + (f" ({data.get('unit')})" if data.get("unit") else "")
            return (title + "\n\n" if data.get("title") else "") + table(head, rows)
        return ""

    # Follow-up suggestions are for the screen, not for paper.
    md = re.sub(r"```\s*followups\s*\n.*?\n\s*```", "", md, flags=re.S | re.I)
    return re.sub(r"```\s*(figures|chart)\s*\n(.*?)\n\s*```", convert, md, flags=re.S | re.I)


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    # A figure written as text in a table ("GHS 9,177,568.91", "73.8%") is a number to
    # whoever opens the sheet; give Excel the number.
    m = re.fullmatch(r"(?:GHS\s*)?(-?[\d,]+(?:\.\d+)?)\s*(%)?", s)
    if m:
        try:
            n = float(m.group(1).replace(",", ""))
            return n / 100 if m.group(2) else n
        except ValueError:
            pass
    return s


def _send_xlsx(title: str, sheets: list, business_code, note: str | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for name, cols, rows in sheets or [("Empty", [], [])]:
        sheet = re.sub(r"[\[\]\*\?/\\:]", " ", str(name))[:28].strip() or "Sheet"
        base, k = sheet, 2
        while sheet in used:
            sheet = f"{base[:25]} {k}"
            k += 1
        used.add(sheet)
        ws = wb.create_sheet(sheet)
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"{_business(business_code)} · figures as at {_as_of()} · exported by "
                   f"{frappe.utils.get_fullname(frappe.session.user)} on "
                   f"{frappe.utils.format_datetime(frappe.utils.now_datetime(), 'd MMM yyyy, HH:mm')}"
                   + (f" · {note}" if note else "")])
        ws.append([])
        ws.append(cols)
        for c in ws[4]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A5"
        for i, col in enumerate(cols, start=1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200] if len(r) >= i])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    buf = io.BytesIO()
    wb.save(buf)
    frappe.local.response.filename = _name(title, "xlsx")
    frappe.local.response.filecontent = buf.getvalue()
    frappe.local.response.type = "binary"


PDF_TEMPLATE = """<!doctype html><html><head><meta charset="utf-8"><style>
body{{font-family:-apple-system,'Segoe UI',Helvetica,Arial,sans-serif;color:#111317;font-size:11.5pt;line-height:1.5}}
.hd{{border-bottom:2px solid #111317;padding-bottom:8px;margin-bottom:18px}}
.hd .b{{font-size:9pt;color:#646973;letter-spacing:.04em;text-transform:uppercase}}
.hd .t{{font-size:10pt;color:#646973;margin-top:2px}}
h1{{font-size:17pt;margin:0 0 12px;letter-spacing:-.01em}}
h2,h3{{font-size:12.5pt;margin:16px 0 6px}}
table{{border-collapse:collapse;width:100%;margin:8px 0 12px;font-size:10pt}}
th{{text-align:left;color:#646973;font-weight:600;border-bottom:1px solid #dfe1e6;padding:6px 8px}}
td{{border-bottom:1px solid #ecedf0;padding:6px 8px}}
ul.checks{{font-family:Menlo,Consolas,monospace;font-size:8pt;color:#646973}}
.ft{{margin-top:26px;border-top:1px solid #dfe1e6;padding-top:8px;font-size:8.5pt;color:#646973}}
</style></head><body>
<div class="hd"><div class="b">{business} · Command Center</div>
<div class="t">Figures as at {as_of} · generated by {who} on {when}</div></div>
<h1>{question}</h1>
{body}
{checks}
<div class="ft">Answered by {model} from {business}'s ERPNext data. The assistant reads the business
within the asking person's permissions and cannot change anything in ERPNext. Projections are
indicative; recommendations are for a person to decide on.</div>
</body></html>"""
